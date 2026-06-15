# -*- coding: utf-8 -*-
"""统计工具函数：数据统计 + Excel写入辅助"""
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULT_DIR = os.path.join(BASE_DIR, 'result')
os.makedirs(RESULT_DIR, exist_ok=True)


# ── 数据统计 ──────────────────────────────────────────────

def zone_stats(data):
    """按专区名称统计，按订单号去重"""
    zones = {}
    for name, group in data.groupby('专区名称'):
        zones[name] = {
            "order_count": int(group['订单号'].nunique()),
            "total_amount": round(float(group['订单金额（元）'].sum()), 2)
        }
    return zones


def supplier_stats(data):
    """按供应商类型统计，按订单号去重"""
    result = {}
    for stype in ["本地供应商", "电商供应商"]:
        stype_data = data[data['供应商类型'] == stype]
        result[stype] = {
            "order_count": int(stype_data['订单号'].nunique()),
            "total_amount": round(float(stype_data['订单金额（元）'].sum()), 2)
        }
    return result


# ── Excel 写入辅助 ────────────────────────────────────────

TITLE_FONT = Font(bold=True, size=18)
HEADER_FONT = Font(bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)


def make_xlsx(suffix, sheet_title):
    """创建 xlsx 工作簿，返回 (wb, ws, xlsx_path)"""
    timestamp = datetime.now().strftime('%Y%m%d%H%M')
    xlsx_path = os.path.join(RESULT_DIR, f'analysis_results_{timestamp}_{suffix}.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    return wb, ws, xlsx_path


def write_title(ws, row, text):
    """写合并居中标题"""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    c = ws.cell(row=row, column=1, value=text)
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal='center')
    return row + 1


def write_header(ws, row, headers=None):
    """写表头行"""
    if headers is None:
        headers = ["专区", "订单数", "销售金额（元）"]
    for col, val in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = HEADER_FONT
        c.border = THIN_BORDER
    return row + 1


def write_row(ws, row, label, count, amount, bold=False):
    """写数据行"""
    c1 = ws.cell(row=row, column=1, value=label)
    c2 = ws.cell(row=row, column=2, value=count)
    c3 = ws.cell(row=row, column=3, value=amount)
    c3.number_format = '#,##0.00'
    for c in (c1, c2, c3):
        c.border = THIN_BORDER
        if bold:
            c.font = Font(bold=True)
    return row + 1


def write_total_row(ws, row, data_start, data_end):
    """写合计行"""
    total_b = sum(ws.cell(row=r, column=2).value or 0 for r in range(data_start, data_end + 1))
    total_c = sum(ws.cell(row=r, column=3).value or 0 for r in range(data_start, data_end + 1))
    ws.cell(row=row, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.cell(row=row, column=2, value=total_b)
    ws.cell(row=row, column=2).border = THIN_BORDER
    ws.cell(row=row, column=2).font = Font(bold=True)
    ws.cell(row=row, column=3, value=round(total_c, 2))
    ws.cell(row=row, column=3).border = THIN_BORDER
    ws.cell(row=row, column=3).font = Font(bold=True)
    ws.cell(row=row, column=3).number_format = '#,##0.00'


def set_col_widths(ws):
    """设置常用列宽"""
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 20


def save_xlsx(wb, xlsx_path):
    """保存并打印路径"""
    wb.save(xlsx_path)
    print(f"Excel已保存: {xlsx_path}")


# ── 区块写入组合 ──────────────────────────────────────────

def write_section_with_supplier(ws, row, title, supplier, zones):
    """写一个区块：标题 + 供应商表 + 专区表（周五/综合/上月共用格式）"""
    row = write_title(ws, row, title)

    # 供应商类型表
    row = write_header(ws, row, ["供应商类型", "订单数", "销售金额（元）"])
    data_start = row
    for stype, s in supplier.items():
        row = write_row(ws, row, stype, s["order_count"], s["total_amount"])
    data_end = row - 1
    write_total_row(ws, row, data_start, data_end)
    row += 1

    # 专区表
    row = write_header(ws, row)
    data_start = row
    for zone, s in zones.items():
        row = write_row(ws, row, zone, s["order_count"], s["total_amount"])
    data_end = row - 1
    write_total_row(ws, row, data_start, data_end)
    return row + 2


def write_section_zone_only(ws, row, title, zones):
    """写一个区块：标题 + 专区表（周一统计用格式）"""
    row = write_title(ws, row, title)
    row = write_header(ws, row)
    data_start = row
    for zone, s in zones.items():
        row = write_row(ws, row, zone, s["order_count"], s["total_amount"])
    data_end = row - 1
    write_total_row(ws, row, data_start, data_end)
    return row + 2
