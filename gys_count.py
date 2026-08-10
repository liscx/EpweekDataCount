# -*- coding: utf-8 -*-
"""本月供应商统计 Top10：采购人数、销售额、订单数"""
import pandas as pd
import os
from datetime import datetime
from openpyxl.styles import Font, Alignment

from stats_utils import (
    make_xlsx, save_xlsx,
    TITLE_FONT, HEADER_FONT, THIN_BORDER,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SOURCE_FILE = os.path.join(BASE_DIR, 'Data', 'source_data.xlsx')


def run():
    """执行供应商统计，返回 (result, xlsx_path)"""
    df = pd.read_excel(SOURCE_FILE)
    df['订单创建时间'] = pd.to_datetime(df['订单创建时间'])

    now = datetime.now()
    current_year = now.year
    current_month = now.month

    # 本月数据
    month_mask = (df['订单创建时间'].dt.year == current_year) & (df['订单创建时间'].dt.month == current_month)
    month_data = df[month_mask].copy()

    # ── 表1：本月供应商服务采购人数量统计 Top10 ──
    supplier_buyer = (
        month_data.groupby('供应商')
        .agg(
            服务采购人数=('采购人', 'nunique'),
            订单总额=('订单金额（元）', 'sum'),
            订单数量=('订单号', 'nunique'),
        )
        .reset_index()
    )
    supplier_buyer['订单总额'] = supplier_buyer['订单总额'].round(2)
    top10_buyer = supplier_buyer.sort_values('服务采购人数', ascending=False).head(10).reset_index(drop=True)

    # ── 表2：本月供应商总销售额统计 Top10 ──
    supplier_amount = (
        month_data.groupby('供应商')
        .agg(
            订单总额=('订单金额（元）', 'sum'),
            订单数量=('订单号', 'nunique'),
        )
        .reset_index()
    )
    supplier_amount['订单总额'] = supplier_amount['订单总额'].round(2)
    top10_amount = supplier_amount.sort_values('订单总额', ascending=False).head(10).reset_index(drop=True)

    # ── 表3：本月供应商总销售订单统计 Top10 ──
    supplier_order = (
        month_data.groupby('供应商')
        .agg(
            订单数量=('订单号', 'nunique'),
            订单总额=('订单金额（元）', 'sum'),
        )
        .reset_index()
    )
    supplier_order['订单总额'] = supplier_order['订单总额'].round(2)
    top10_order = supplier_order.sort_values('订单数量', ascending=False).head(10).reset_index(drop=True)

    month_label = f"{current_year}年{current_month}月"

    # 控制台输出
    print(f"\n{month_label}供应商统计 Top10")
    print("\n1. 服务采购人数量统计:")
    print(top10_buyer.to_string(index=False))
    print("\n2. 总销售额统计:")
    print(top10_amount.to_string(index=False))
    print("\n3. 总销售订单统计:")
    print(top10_order.to_string(index=False))

    # ── 导出 Excel（三张表放在同一个 sheet）──
    wb, ws, xlsx_path = make_xlsx("GYS", "供应商统计")

    row = 1

    # ── 表1：采购人数量 ──
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    c = ws.cell(row=row, column=1, value=f"{month_label}供应商服务采购人数量统计 Top10")
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal='center')
    row += 1

    headers1 = ["供应商", "服务采购人数", "订单总额（元）", "订单数量"]
    for col, val in enumerate(headers1, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = HEADER_FONT
        c.border = THIN_BORDER
    row += 1

    for _, r in top10_buyer.iterrows():
        vals = [r['供应商'], int(r['服务采购人数']), r['订单总额'], int(r['订单数量'])]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = THIN_BORDER
            if isinstance(val, float):
                c.number_format = '#,##0.00'
        row += 1

    row += 2  # 间隔

    # ── 表2：销售额 ──
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    c = ws.cell(row=row, column=1, value=f"{month_label}供应商总销售额统计 Top10")
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal='center')
    row += 1

    headers2 = ["供应商", "订单总额（元）", "订单数量"]
    for col, val in enumerate(headers2, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = HEADER_FONT
        c.border = THIN_BORDER
    row += 1

    for _, r in top10_amount.iterrows():
        vals = [r['供应商'], r['订单总额'], int(r['订单数量'])]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = THIN_BORDER
            if isinstance(val, float):
                c.number_format = '#,##0.00'
        row += 1

    row += 2  # 间隔

    # ── 表3：订单数 ──
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    c = ws.cell(row=row, column=1, value=f"{month_label}供应商总销售订单统计 Top10")
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal='center')
    row += 1

    headers3 = ["供应商", "订单数量", "订单总额（元）"]
    for col, val in enumerate(headers3, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = HEADER_FONT
        c.border = THIN_BORDER
    row += 1

    for _, r in top10_order.iterrows():
        vals = [r['供应商'], int(r['订单数量']), r['订单总额']]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = THIN_BORDER
            if isinstance(val, float):
                c.number_format = '#,##0.00'
        row += 1

    # 列宽
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 14

    save_xlsx(wb, xlsx_path)

    result = {
        "type": "gys_top10",
        "month": f"{current_year}-{current_month:02d}",
        "top10_buyer": top10_buyer.to_dict('records'),
        "top10_amount": top10_amount.to_dict('records'),
        "top10_order": top10_order.to_dict('records'),
    }
    return result, xlsx_path


if __name__ == "__main__":
    run()
