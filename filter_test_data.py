# -*- coding: utf-8 -*-
"""过滤测试订单数据：将测试数据移到单独的sheet，返回干净数据

两类过滤规则：
1. 关键词过滤（is_test_order）：采购企业/采购部门/收货地址 含"测试"等字样
2. 订单号精确过滤：订单号命中 EXCLUDE_FILE 清单 → 整单（所有明细行）剔除

命中结果：
- 关键词命中的行  → "测试数据" sheet
- 清单命中的订单  → "黑名单订单" sheet
- 其余行保留在 Sheet1，作为后续统计的干净数据
"""
import pandas as pd
import os
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SOURCE_FILE = os.path.join(BASE_DIR, 'Data', 'source_data.xlsx')
# 手动排除的订单号清单：每行一个订单号，# 开头为注释（见同目录 exclude_order_ids.txt）
EXCLUDE_FILE = os.path.join(BASE_DIR, 'exclude_order_ids.txt')
ORDER_COL = '订单号'


def is_test_order(row):
    """判断是否为测试订单"""
    # 采购企业包含"测试"
    if pd.notna(row.get('采购企业')) and '测试' in str(row['采购企业']):
        return True

    # 采购部门包含"测试"
    if pd.notna(row.get('采购部门')) and '测试' in str(row['采购部门']):
        return True

    # 采购部门为"系统管理部"
    if pd.notna(row.get('采购部门')) and str(row['采购部门']).strip() == '系统管理部':
        return True

    # 收货地址包含"测试"、"国泰测试"、"测试地址"
    if pd.notna(row.get('收货地址')):
        addr = str(row['收货地址'])
        if '测试' in addr or '国泰测试' in addr or '测试地址' in addr:
            return True

    return False


def load_exclude_order_ids():
    """读取排除订单号清单文件，返回 set；文件不存在时返回空集"""
    ids = set()
    if not os.path.exists(EXCLUDE_FILE):
        return ids
    with open(EXCLUDE_FILE, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            ids.add(line)
    return ids


def _norm_order_id(v):
    """归一化订单号用于精确比较：去掉首尾空格；避免 Excel 数字型订单号读出 12345.0"""
    if pd.isna(v):
        return ''
    s = str(v).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]
    return s


def is_blacklisted_order(row, exclude_ids):
    """订单号精确匹配：命中清单即视为需整单剔除"""
    return _norm_order_id(row.get(ORDER_COL)) in exclude_ids


def _write_df_to_sheet(ws, df):
    """将 DataFrame（含表头）写入 worksheet"""
    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def filter_test_data():
    """过滤测试数据：Sheet1 保留干净数据，关键词命中/清单命中分别写入独立sheet"""
    print("=== 过滤测试订单 ===")

    # 读取源数据（第一个 sheet）
    df = pd.read_excel(SOURCE_FILE)
    total_before = len(df)
    print(f"原始数据行数: {total_before}")

    # 规则1：关键词标记测试订单
    kw_mask = df.apply(is_test_order, axis=1)

    # 规则2：订单号精确过滤（整单剔除）
    exclude_ids = load_exclude_order_ids()
    if exclude_ids:
        print(f"排除清单订单号数: {len(exclude_ids)}")
        bl_mask = df[ORDER_COL].map(lambda v: _norm_order_id(v) in exclude_ids)
    else:
        bl_mask = pd.Series(False, index=df.index)

    black_data = df[bl_mask]
    kw_test_data = df[kw_mask & ~bl_mask]
    clean_data = df[~(kw_mask | bl_mask)]

    test_count = len(kw_test_data)
    black_count = len(black_data)
    clean_count = len(clean_data)
    print(f"关键词测试订单行数: {test_count}")
    print(f"清单命中订单行数(整单剔除): {black_count}")
    print(f"有效订单行数: {clean_count}")

    # 创建新的工作簿
    wb = Workbook()

    # 干净数据写入 Sheet1
    ws_clean = wb.active
    ws_clean.title = 'Sheet1'
    _write_df_to_sheet(ws_clean, clean_data)

    # 关键词命中数据写入 "测试数据" sheet
    if test_count > 0:
        ws_test = wb.create_sheet('测试数据')
        _write_df_to_sheet(ws_test, kw_test_data)
        print(f"关键词命中数据已保存到 '测试数据' sheet")

    # 订单号清单命中数据写入 "黑名单订单" sheet
    if black_count > 0:
        ws_black = wb.create_sheet('黑名单订单')
        _write_df_to_sheet(ws_black, black_data)
        print(f"清单命中数据已保存到 '黑名单订单' sheet")

    # 保存文件
    wb.save(SOURCE_FILE)
    print(f"source_data.xlsx 已更新，仅保留有效数据")

    return clean_data


if __name__ == "__main__":
    filter_test_data()
